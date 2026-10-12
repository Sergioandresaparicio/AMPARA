"""
Pruebas Unitarias para Tool 4: Petition Generator (AMPARA / Hermes)
Archivo: agent/tests/test_petition_engine.py
"""

import unittest
from agent.tools.petition_engine import PetitionGeneratorEngine, draft_template_engine

class TestPetitionGenerator(unittest.TestCase):

    def setUp(self):
        self.engine = PetitionGeneratorEngine()

    def test_generate_tutela_con_medida_provisional(self):
        case_data = {
            "patient_name": "CARMEN ROSA SUÁREZ",
            "patient_id": "39.754.210",
            "eps_name": "EPS SANITAS",
            "diagnosis": "C50.9 Cáncer de Mama Estadio III",
            "urgency_level": "critical",
            "critical_escalation": True,
            "facts": [{"details": "La EPS no entrega medicamento oncológico Trastuzumab"}],
            "city": "Medellín"
        }
        res = draft_template_engine(case_data)
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["action_type"], "TUTELA_CON_MEDIDA_PROVISIONAL")
        self.assertEqual(res["subject_matter"], "MEDICAMENTO")
        self.assertTrue(res["metadata"]["has_provisional_measures"])
        self.assertTrue(res["metadata"]["human_review_required"])
        # Verificar contenido procesal
        md = res["generated_draft_markdown"]
        self.assertIn("MEDIDA CAUTELAR PROVISIONAL", md)
        self.assertIn("ARTÍCULO 7 DEL DECRETO 2591 DE 1991", md)
        self.assertIn("CARMEN ROSA SUÁREZ", md)
        self.assertIn("EPS SANITAS", md)
        self.assertIn("JURAMENTO DE NO DUPLICIDAD", md)
        self.assertIn("Artículo 37 del Decreto 2591 de 1991", md)

    def test_generate_derecho_peticion(self):
        case_data = {
            "patient_name": "JORGE ENRIQUE GÓMEZ",
            "patient_id": "19.345.890",
            "eps_name": "EPS SURA",
            "phase": "peticion_previa",
            "prescription_item": "Insulina Glargina 100 UI",
            "city": "Cali"
        }
        res = draft_template_engine(case_data, template_type="DERECHO_PETICION")
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["action_type"], "DERECHO_PETICION")
        md = res["generated_draft_markdown"]
        self.assertIn("DERECHO DE PETICIÓN EN SALUD", md)
        self.assertIn("ARTÍCULO 23 CONSTITUCIÓN POLÍTICA", md)
        self.assertIn("Insulina Glargina 100 UI", md)

    def test_generate_incidente_desacato(self):
        case_data = {
            "patient_name": "MARÍA ELENA PÉREZ",
            "patient_id": "52.458.912",
            "eps_name": "COMPENSAR EPS",
            "phase": "desacato_post_fallo",
            "city": "Bogotá D.C."
        }
        res = draft_template_engine(case_data)
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["action_type"], "INCIDENTE_DESACATO")
        md = res["generated_draft_markdown"]
        self.assertIn("TRÁMITE INCIDENTAL DE DESACATO", md)
        self.assertIn("ARTÍCULOS 27 Y 52 DEL DECRETO 2591 DE 1991", md)

    def test_integration_with_doc_parser_and_rag(self):
        case_data = {
            "patient_name": "LUCÍA FERNÁNDEZ",
            "patient_id": "41.890.123",
            "eps_name": "NUEVA EPS",
            "diagnosis": "Insuficiencia Renal Crónica N18.5",
            "urgency_level": "critical",
            "city": "Barranquilla"
        }
        # Simular salida del Document Analyzer (Tool 3)
        doc_analyses = [
            {
                "document_type": "FORMULA_MEDICA",
                "document_description": "Fórmula médica de especialista",
                "summary_for_tutela": "Fórmula médica expedida donde se prescribió con carácter necesario: Eritropoyetina 4000 UI.",
                "extracted_entities": {
                    "prescribed_items": ["Eritropoyetina 4000 UI"]
                },
                "file_info": {"file_name": "formula_nefrologia.pdf"}
            }
        ]
        # Simular salida del Legal RAG (Tool 2)
        legal_sources = [
            {
                "chunk_id": "DEC-2591-1991-ART-7",
                "title": "Decreto 2591 de 1991 - Artículo 7",
                "text": "Desde la presentación de la solicitud, cuando el juez expresamente lo considere necesario y urgente para proteger el derecho, suspenderá la aplicación del acto objeto de la tutela."
            }
        ]

        res = draft_template_engine(
            case_data=case_data,
            legal_sources=legal_sources,
            document_analyses=doc_analyses
        )
        self.assertEqual(res["status"], "success")
        md = res["generated_draft_markdown"]
        # Debe contener el ítem extraído del doc_parser
        self.assertIn("Eritropoyetina 4000 UI", md)
        # Debe contener la cita del RAG
        self.assertIn("DEC-2591-1991-ART-7", md)
        # Debe contener el anexo documental
        self.assertIn("formula_nefrologia.pdf", md)
        # Trazabilidad S4
        self.assertTrue(len(res["sources_for_s4"]) >= 1)
        self.assertEqual(res["sources_for_s4"][0]["id"], "DEC-2591-1991-ART-7")

    def test_different_health_subject_matters(self):
        # Caso de viáticos y transporte
        case_transporte = {
            "patient_name": "CARLOS MEJÍA",
            "patient_id": "10.234.567",
            "eps_name": "COOSALUD",
            "facts": [{"details": "Requiere transporte intermunicipal y viáticos de alojamiento en otra ciudad"}]
        }
        res_transp = draft_template_engine(case_transporte)
        self.assertEqual(res_transp["subject_matter"], "TRANSPORTE_VIATICOS")

        # Caso de insumos (silla de ruedas)
        case_insumo = {
            "patient_name": "ANA SOFÍA TORRES",
            "patient_id": "23.456.789",
            "eps_name": "FAMISANAR",
            "facts": [{"details": "Solicitud de silla de ruedas neurológica y pañales desechables"}]
        }
        res_insumo = draft_template_engine(case_insumo)
        self.assertEqual(res_insumo["subject_matter"], "INSUMOS_DISPOSITIVOS")

    def test_export_to_docx_direct_and_integrated(self):
        import os
        import tempfile
        from docx import Document
        from agent.tools.petition_engine import export_petition_to_docx

        # 1. Prueba directa export_petition_to_docx
        md_sample = "# ACCIÓN DE TUTELA\n\nYo, **CARLOS PÉREZ**, interpongo tutela contra **EPS SURA**.\n\n- Anexo 1: Historia clínica"
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            out_file = export_petition_to_docx(
                document_markdown=md_sample,
                output_path=tmp_path,
                case_data={"patient_id": "12345"},
                action_type="TUTELA"
            )
            self.assertTrue(os.path.exists(out_file))
            self.assertGreater(os.path.getsize(out_file), 1000)

            # Verificar que python-docx abre el archivo generado
            doc = Document(out_file)
            full_text = " ".join([p.text for p in doc.paragraphs])
            self.assertIn("ACCIÓN DE TUTELA", full_text)
            self.assertIn("CARLOS PÉREZ", full_text)
            self.assertIn("EPS SURA", full_text)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

        # 2. Prueba integrada vía draft_template_engine(export_docx=True)
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp2:
            tmp_path2 = tmp2.name

        try:
            case_data = {
                "patient_name": "MARÍA RESTREPO",
                "patient_id": "60.123.456",
                "eps_name": "SANITAS",
                "diagnosis": "Cardiopatía",
                "facts": [{"details": "Demora en cita con cardiología"}]
            }
            res = draft_template_engine(
                case_data=case_data,
                export_docx=True,
                output_docx_path=tmp_path2
            )
            self.assertEqual(res["status"], "success")
            self.assertIn("docx_file_path", res)
            self.assertTrue(os.path.exists(res["docx_file_path"]))
            self.assertGreater(os.path.getsize(res["docx_file_path"]), 1000)

            doc2 = Document(res["docx_file_path"])
            full_text2 = " ".join([p.text for p in doc2.paragraphs])
            self.assertIn("MARÍA RESTREPO", full_text2)
            self.assertIn("SANITAS", full_text2)
        finally:
            if os.path.exists(tmp_path2):
                os.remove(tmp_path2)

    def test_first_person_drafting_for_propio(self):
        """Verifica que a título propio el memorial se redacte en primera persona del singular."""
        case_propio = {
            "patient_name": "CARMEN ROSA SUÁREZ",
            "patient_id": "39.754.210",
            "eps_name": "EPS SANITAS",
            "diagnosis": "Cáncer de Mama",
            "prescription_item": "Trastuzumab 600 mg",
            "acting_as": "propio",
            "city": "Medellín"
        }
        res = draft_template_engine(case_propio)
        md = res["generated_draft_markdown"]
        
        # Debe contener redacción en primera persona
        self.assertIn("actuando en mi propio nombre y representación", md)
        self.assertIn("Me encuentro formalmente afiliado(a) a la EPS accionada", md)
        self.assertIn("Fui diagnosticado(a) clínicamente", md)
        self.assertIn("mi médico tratante adscrito", md)
        self.assertIn("TUTELAR y AMPARAR de manera inmediata mis derechos fundamentales", md)
        self.assertIn("al suscrito accionante **CARMEN ROSA SUÁREZ**", md)
        self.assertIn("Bajo la gravedad del juramento manifiesto a usted, señor Juez, que no he interpuesto otra acción", md)

    def test_third_person_drafting_for_agente_oficioso(self):
        """Verifica que como agente oficioso/abogado se redacte en tercera persona con sustento en Art. 10 Dec. 2591/91."""
        case_agente = {
            "patient_name": "CARLOS ALBERTO MENDOZA",
            "patient_id": "19.876.543",
            "eps_name": "NUEVA EPS",
            "diagnosis": "Insuficiencia Cardíaca Terminal",
            "prescription_item": "Marcapasos Bicameral",
            "acting_as": "agente_oficioso",
            "representative_name": "DR. JUAN PABLO GIRALDO",
            "representative_id": "79.123.456",
            "city": "Bogotá D.C."
        }
        res = draft_template_engine(case_agente)
        md = res["generated_draft_markdown"]
        
        # Encabezado
        self.assertIn("DR. JUAN PABLO GIRALDO (COMO AGENTE OFICIOSO DE CARLOS ALBERTO MENDOZA)", md)
        # Personería e invocación legal
        self.assertIn("actuando en calidad de **AGENTE OFICIOSO** conforme a lo consagrado en el **Artículo 10 del Decreto 2591 de 1991**", md)
        self.assertIn("en defensa y amparo de los derechos fundamentales de **CARLOS ALBERTO MENDOZA**", md)
        # Hechos canónicos en tercera persona
        self.assertIn("Mi agenciado(a), **CARLOS ALBERTO MENDOZA**, se encuentra formalmente afiliado(a)", md)
        self.assertIn("Mi agenciado(a) padece de: **Insuficiencia Cardíaca Terminal**", md)
        self.assertIn("Marcapasos Bicameral", md)
        self.assertIn("el estado clínico de mi agenciado(a)", md)
        # Pretensiones en tercera persona
        self.assertIn("de mi agenciado(a) **CARLOS ALBERTO MENDOZA**", md)
        self.assertIn("entregar efectivamente a mi agenciado(a) **CARLOS ALBERTO MENDOZA**", md)
        # Juramento y firma
        self.assertIn("ni mi agenciado(a) ni el suscrito agente oficioso hemos interpuesto otra acción", md)
        self.assertIn("Agente Oficioso de **CARLOS ALBERTO MENDOZA**", md)

if __name__ == "__main__":
    unittest.main()
